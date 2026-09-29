"""Lossless XLSX intake and normalization of the supplied three-row EKP format.

This is a classifier, not a dispatcher: conditions describe the source matrix.
Unresolved abbreviations and scenario codes never become invented routing rules.
"""
import hashlib
import io
import re
import zipfile
from collections import Counter
from datetime import date, datetime
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MAX_UPLOAD = 10 * 1024 * 1024
PARSER_VERSION = 1


class ClassifierFormatError(ValueError):
    pass


def text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def canonical(value):
    return re.sub(r"\s+", " ", text(value)).casefold().replace("ё", "е")


def json_value(value):
    return value.isoformat() if isinstance(value, (datetime, date)) else value


def issue(code, message, sheet=None, row=None, cell=None):
    return {k: v for k, v in dict(code=code, message=message, sheet=sheet, row=row, cell=cell).items() if v is not None}


def condition_for(label):
    label = canonical(label)
    known = {
        "служба 101 (признак нд - нет доступа не выбран)": ("no_access", False),
        "служба 101 (выбран признак нд - нет доступа)": ("no_access", True),
        "одс псц (выбран признак ул - угроза людям)": ("threat_to_people", True),
        "одс псц (выбран признак пп - пострадавшие погибшие)": ("casualties", True),
        "одс псц (выбран признак нд - нет доступа)": ("no_access", True),
        "выбран признак правонарушение": ("offense", True),
        "выбран признак пострадавшие": ("casualties", True),
        "классификатор смп (признак пострадавшие не выбран)": ("casualties", False),
        "классификатор смп (выбран признак пострадавшие)": ("casualties", True),
        "классификатор смп (выбран признак пострадавшие не на месте)": ("casualties_not_on_scene", True),
        "газификация": ("gasified", True),
        "угроза людям": ("threat_to_people", True),
        "пострадавшие/погибшие": ("casualties", True),
        "мед. помощь": ("medical_help", True),
        "треб. эвакуация": ("evacuation", True),
        "постр / погибшие": ("casualties", True),
        "перекрытие движение": ("road_blocked", True),
        "тоннель": ("tunnel", True),
        "на объектах связи": ("communications_facility", True),
        "стройка": ("construction", True),
    }
    if not label or label == "реагирование всегда":
        return {"kind": "always", "review_required": False}
    if label in known:
        feature, value = known[label]
        return {"kind": "flag", "feature": feature, "value": value, "review_required": False}
    if label in ("признак не выбран", "признаки не выбраны", "одс псц (другие признаки не выбраны)"):
        return {"kind": "default", "review_required": False}
    # The exact combination/meaning of abbreviations is not specified by the file.
    return {"kind": "unresolved", "source_text": label, "review_required": True}


def header_columns(ws, merged_values):
    max_col = ws.max_column
    columns = []
    for col in range(1, max_col + 1):
        levels = [text(merged_values.get((r, col), ws.cell(r, col).value)) for r in range(1, 4)]
        path = list(dict.fromkeys(x for x in levels if x))
        item = {"column": get_column_letter(col), "index": col, "levels": levels, "path": path}
        if col >= 15:
            group = levels[0]
            # MCHS contains distinct service recipients; other second-level names may be delivery channels.
            service = levels[1] if canonical(group) == "классификатор мчс" else group
            channel_groups = {"департамент рбипк (гку мосбез)", "дгп (департамент градостроительной политики)", "гуп мср"}
            channel = next((v for v in reversed(levels[1:]) if v and v != group), "") if canonical(group) in channel_groups else ""
            if channel:
                condition_label = ""
            elif canonical(group) == "классификатор мчс":
                condition_label = levels[2] if levels[2] != levels[1] else ""
            else:
                condition_label = next((v for v in reversed(levels[1:]) if v and v != group), "")
            service = service or f"Неизвестная служба ({item['column']})"
            key = hashlib.sha256(canonical(service).encode()).hexdigest()[:16]
            predicate = condition_for(condition_label)
            if not group:
                predicate = {"kind": "unresolved", "source_text": "Отсутствует заголовок службы", "review_required": True}
            item.update(service_key=key, service_name=service, channel=channel, condition_label=condition_label, condition=predicate)
        columns.append(item)
    # A generic 'not selected' means none of this service's explicit options.
    for column in columns[14:]:
        if column["condition"]["kind"] == "default":
            peers = [c["condition"] for c in columns[14:] if c["service_key"] == column["service_key"] and c is not column]
            features = sorted({p["feature"] for p in peers if p["kind"] == "flag"})
            unresolved = any(p["review_required"] for p in peers) or not features
            column["condition"] = {"kind": "all_false", "features": features, "review_required": unresolved}
    return columns


def parse_workbook(data, filename=""):
    if len(data) > MAX_UPLOAD:
        raise ClassifierFormatError("Размер XLSX превышает 10 МБ")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                raise ClassifierFormatError("Распакованный XLSX превышает 100 МБ")
        wb = load_workbook(io.BytesIO(data), data_only=False)
        cached = load_workbook(io.BytesIO(data), data_only=True)
    except ClassifierFormatError:
        raise
    except Exception as exc:
        raise ClassifierFormatError("Повреждённый XLSX") from exc
    result = {"format": "generic", "source_sha256": hashlib.sha256(data).hexdigest(), "rows": [],
              "validation_errors": [], "warnings": [],
              "metadata": {"parser_version": PARSER_VERSION, "sheets": [], "services": [], "statistics": {}}}
    services = {}
    seen_codes = {}
    modes = set()
    try:
        if sum(s.max_row * s.max_column for s in wb) > 1_000_000:
            raise ClassifierFormatError("Слишком большая область данных XLSX (максимум 1 млн ячеек)")
        for ws in wb:
            if ws.max_row > 10000 or ws.max_column > 256:
                raise ClassifierFormatError("Лист превышает 10000 строк или 256 столбцов")
            values = cached[ws.title]
            max_row, max_col = ws.max_row, ws.max_column
            is_ekp = "екп" in canonical(ws["L2"].value) or canonical(ws["K2"].value) == "итоговый тип происшествия"
            modes.add("ekp" if is_ekp else "generic")
            merged = {}
            for area in ws.merged_cells.ranges:
                for row in range(area.min_row, area.max_row + 1):
                    for col in range(area.min_col, area.max_col + 1):
                        merged[row, col] = values.cell(area.min_row, area.min_col).value
            if not is_ekp:
                headers = [text(c.value) or f"column_{c.column}" for c in ws[1]]
                duplicate = [k for k, v in Counter(headers).items() if v > 1]
                if duplicate:
                    result["warnings"].append(issue("DUPLICATE_HEADERS", "Повторяющиеся заголовки сохранены с координатами", ws.title))
                for row in values.iter_rows(min_row=2):
                    if not any(c.value is not None for c in row):
                        continue
                    names = [f"{h} [{get_column_letter(i+1)}]" if h in duplicate else h for i, h in enumerate(headers)]
                    record = {names[i]: json_value(c.value) for i, c in enumerate(row)}
                    record.update(__sheet__=ws.title, __row__=row[0].row)
                    result["rows"].append(record)
                result["metadata"]["sheets"].append({"name": ws.title, "format": "generic", "columns": ws.max_column, "rows": ws.max_row})
                continue
            if ws.max_column < 90:
                result["validation_errors"].append(issue("TRUNCATED_EKP", "Ожидается полный формат ЕКП A:CL (90 столбцов)", ws.title))
            for coord, label in {"E2": "Номер", "G2": "112 - Признак.1 (тип происшествия)", "K2": "Итоговый тип происшествия", "M1": "Сценарий реагирования", "N1": "Главная служба"}.items():
                if canonical(ws[coord].value) != canonical(label):
                    result["validation_errors"].append(issue("HEADER_MISMATCH", f"Ожидается заголовок «{label}»", ws.title, cell=coord))
            columns = header_columns(ws, merged)
            result["metadata"]["sheets"].append({"name": ws.title, "format": "ekp", "rows": ws.max_row, "column_count": ws.max_column,
                "header_rows": 3, "columns": columns, "merged_ranges": [str(m) for m in ws.merged_cells.ranges]})
            for column in columns[14:]:
                services[column["service_key"]] = {"key": column["service_key"], "name": column["service_name"]}
                if column["condition"]["review_required"]:
                    result["warnings"].append(issue("UNRESOLVED_CONDITION", f"Нужно уточнить условие: {column['condition_label'] or column['service_name']}", ws.title, cell=column["column"]+"3"))
            group_name = None
            subgroup = None
            for row_no in range(4, max_row + 1):
                raw = {get_column_letter(c): json_value(values.cell(row_no, c).value) for c in range(1, max_col + 1)}
                if not any(v is not None for v in raw.values()):
                    continue
                formulas = {get_column_letter(c): ws.cell(row_no, c).value for c in range(1, max_col + 1) if ws.cell(row_no, c).data_type == "f"}
                for col in formulas:
                    if raw[col] is None:
                        result["validation_errors"].append(issue("FORMULA_NO_CACHE", "Формула не имеет рассчитанного значения; пересохраните файл в Excel", ws.title, row_no, f"{col}{row_no}"))
                resolved = {get_column_letter(c): merged.get((row_no, c), values.cell(row_no, c).value) for c in range(1, max_col + 1)}
                code, title = text(resolved.get("E")), text(resolved.get("K"))
                root_heading = not title and bool(text(resolved.get("F"))) and not any(text(resolved.get(c)) for c in ("A","B","C","D","G","H","I"))
                if root_heading:
                    group_name = text(resolved["F"])
                    subgroup = None
                record = {"source_sheet": ws.title, "source_row": row_no, "raw": raw, "formulas": formulas, "kind": "group" if root_heading else "incident", "warnings": []}
                if root_heading:
                    record["group_name"] = group_name
                    result["rows"].append(record)
                    continue
                if not code or not title:
                    record["kind"] = "unrecognized"
                    result["validation_errors"].append(issue("INCOMPLETE_INCIDENT", "Нет номера или итогового типа происшествия", ws.title, row_no))
                if len(code) > 100 or len(title) > 1000:
                    result["validation_errors"].append(issue("VALUE_TOO_LONG", "Слишком длинный код или название", ws.title, row_no))
                if code in seen_codes:
                    result["validation_errors"].append(issue("DUPLICATE_CODE", f"Код {code} повторяет {seen_codes[code]}", ws.title, row_no, f"E{row_no}"))
                seen_codes[code] = f"{ws.title}!E{row_no}"
                if text(resolved.get("F")):
                    subgroup = text(resolved["F"])
                scenario = text(resolved.get("M")) or None
                if not scenario or canonical(scenario) == "без сценария":
                    warning = issue("MISSING_SCENARIO", "Сценарий реагирования не задан; автоматическая маршрутизация по его коду недоступна", ws.title, row_no, f"M{row_no}")
                    record["warnings"].append(warning)
                    result["warnings"].append(warning)
                rules = []
                for column in columns[14:]:
                    value = text(resolved.get(column["column"]))
                    if not value:
                        continue
                    effect = "NO_RESPONSE" if canonical(value) == "нет реагирования" else "CLASSIFICATION"
                    rules.append({"source_cell": f"{column['column']}{row_no}", "column": column["column"],
                        "service_key": column["service_key"], "service_name": column["service_name"], "channel": column["channel"],
                        "classification": value, "effect": effect, "condition": column["condition"], "condition_label": column["condition_label"]})
                record.update(code=code, title=title, group_name=group_name, subgroup=subgroup,
                    features={c: text(resolved.get(c)) or None for c in ("G","H","I","J")},
                    ekp_type=text(resolved.get("L")) or None, response_scenario_code=scenario,
                    lead_service_code=text(resolved.get("N")) or None, rules=rules,
                    review_required=bool(record["warnings"]) or any(r["condition"]["review_required"] for r in rules))
                result["rows"].append(record)
        if "ekp" in modes:
            result["format"] = "ekp"
            if "generic" in modes:
                result["validation_errors"].append(issue("MIXED_FORMATS", "В книге ЕКП есть листы другого формата; загрузите их отдельно"))
        if not result["rows"]:
            result["validation_errors"].append(issue("EMPTY_FILE", "Файл не содержит строк данных"))
        incidents = [r for r in result["rows"] if r.get("kind") == "incident"]
        if result["format"] == "ekp" and not incidents:
            result["validation_errors"].append(issue("NO_INCIDENTS", "В файле нет происшествий"))
        if result["format"] == "ekp" and "искл" in canonical(filename) and any("пожар" in canonical(r.get("title")) or "задымлен" in canonical(r.get("title")) for r in incidents):
            result["warnings"].append(issue("FILENAME_EXCLUSION", "В имени файла указано исключение, но пожарные строки присутствуют. Все строки сохранены, автоматического исключения нет."))
        result["metadata"]["services"] = sorted(services.values(), key=lambda x: x["name"])
        result["metadata"]["statistics"] = {
            "source_rows": len(result["rows"]), "incident_count": len(incidents),
            "group_count": sum(r.get("kind") == "group" for r in result["rows"]),
            "service_count": len(services), "rule_count": sum(len(r.get("rules", [])) for r in incidents),
            "no_response_rule_count": sum(rule["effect"] == "NO_RESPONSE" for r in incidents for rule in r["rules"]),
            "review_required_count": sum(r.get("review_required", False) for r in incidents),
            "warning_count": len(result["warnings"]), "error_count": len(result["validation_errors"])}
        return result
    finally:
        wb.close()
        cached.close()
