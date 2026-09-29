"""Classifier parsing tests. Fixtures exercise EKP shape, not production data copies."""
import io
import os
from pathlib import Path
import pytest
from openpyxl import Workbook, load_workbook
from app.services.classifier import parse_workbook, ClassifierFormatError

def sample_book():
    wb = Workbook()
    ws = wb.active
    ws.title = "ЕКП"
    for coord, value in {
        "E2":"Номер", "G2":"112 - Признак.1 (тип происшествия)", "K2":"Итоговый тип происшествия",
        "L2":"ТИП происшествия ЕКП 35", "M1":"Сценарий реагирования", "N1":"Главная служба",
        "Y1":"Классификатор СМП", "Y2":"Классификатор СМП (признак Пострадавшие не выбран)",
        "Z2":"Классификатор СМП (выбран признак Пострадавшие)",
        "AA2":"Классификатор СМП (выбран признак Пострадавшие не на месте)",
        "AM1":"Мосгортранс", "AM3":"признак не выбран", "AN3":"постр / погибшие", "AO3":"перекрытие движение",
        "AY1":"Мосводоканал", "CK1":"ДГП (Департамент градостроительной политики)",
        "CK2":"интеграция", "CL2":"АРМ-112",
        "AI1":"Классификатор ФСБ", "AI3":"признак не выбран", "AJ3":">5 чел / ОД",
        "E4":2, "F4":"Транспорт", "E5":"002010000", "G5":"ДТП", "K5":"ДТП учебное",
        "M5":"2_4", "N5":"Police", "Y5":"нет реагирования", "Z5":"АВТО",
        "AO5":"карточка-112", "AY5":"карточка-112", "CL5":"карточка-112",
    }.items():
        ws[coord] = value
    for area in ("Y1:AA1", "AM1:AO1", "CK1:CL1", "AI1:AJ1"):
        ws.merge_cells(area)
    return wb

def binary(wb):
    stream=io.BytesIO()
    wb.save(stream)
    return stream.getvalue()

def test_full_width_headers_rules_and_leading_zero():
    parsed=parse_workbook(binary(sample_book()))
    assert not parsed["validation_errors"]
    assert len(parsed["rows"])==2
    row=parsed["rows"][1]
    assert row["code"]=="002010000" and row["group_name"]=="Транспорт"
    assert len(row["raw"])==90 and row["raw"]["AY"]=="карточка-112" and row["raw"]["CL"]=="карточка-112"
    rules={r["column"]:r for r in row["rules"]}
    assert rules["Y"]["effect"]=="NO_RESPONSE"
    assert rules["Y"]["condition"]=={"kind":"flag","feature":"casualties","value":False,"review_required":False}
    assert rules["Z"]["condition"]["value"] is True
    assert rules["AO"]["condition"]["feature"]=="road_blocked"
    assert rules["CL"]["channel"]=="АРМ-112"
    assert rules["CL"]["condition"]["kind"]=="always"

def test_unknown_conditions_stay_unresolved():
    parsed=parse_workbook(binary(sample_book()))
    columns={r["column"]:r for r in parsed["metadata"]["sheets"][0]["columns"]}
    assert columns["AJ"]["condition"]["kind"]=="unresolved"
    assert columns["AI"]["condition"]["review_required"]
    assert columns["AM"]["condition"]["features"]==["casualties","road_blocked"]
    assert not columns["AM"]["condition"]["review_required"]

def test_missing_scenario_is_warning_and_not_invented():
    wb=sample_book();wb.active["M5"]=None
    parsed=parse_workbook(binary(wb))
    assert not parsed["validation_errors"]
    assert parsed["rows"][1]["response_scenario_code"] is None
    assert parsed["rows"][1]["review_required"]
    assert any(w["code"]=="MISSING_SCENARIO" for w in parsed["warnings"])

def test_duplicate_codes_block_commit_and_preserve_rows():
    wb=sample_book();ws=wb.active;ws["E6"]="002010000";ws["K6"]="Дубликат"
    parsed=parse_workbook(binary(wb))
    assert len(parsed["rows"])==3
    assert any(e["code"]=="DUPLICATE_CODE" for e in parsed["validation_errors"])

def test_formula_without_cached_value_is_not_silently_dropped():
    wb=sample_book();wb.active["E5"]='=1+2'
    parsed=parse_workbook(binary(wb))
    assert parsed["rows"][1]["formulas"]["E"]=="=1+2"
    assert any(e["code"]=="FORMULA_NO_CACHE" for e in parsed["validation_errors"])

def test_generic_import_preserves_all_columns_and_duplicate_headers():
    wb=Workbook();ws=wb.active
    ws.append(["Название"]*90);ws.append(list(range(90)))
    parsed=parse_workbook(binary(wb))
    assert parsed["format"]=="generic"
    assert parsed["rows"][0]["Название [CL]"]==89
    assert len(parsed["rows"][0])==92

def test_invalid_file():
    with pytest.raises(ClassifierFormatError,match="Повреждённый"):
        parse_workbook(b"invalid")

def test_empty_file():
    assert parse_workbook(binary(Workbook()))["validation_errors"][0]["code"]=="EMPTY_FILE"

def test_filename_never_removes_fire_rows():
    wb=sample_book();wb.active["K5"]="пожар"
    parsed=parse_workbook(binary(wb),"искл_пожар.xlsx")
    assert len([r for r in parsed["rows"] if r["kind"]=="incident"])==1
    assert any(w["code"]=="FILENAME_EXCLUSION" for w in parsed["warnings"])

def test_changed_schema_and_mixed_sheet_are_rejected():
    wb=sample_book();wb.active["E2"]="Не номер";wb.create_sheet("Посторонний")
    parsed=parse_workbook(binary(wb))
    assert {"HEADER_MISMATCH","MIXED_FORMATS"} <= {e["code"] for e in parsed["validation_errors"]}

@pytest.mark.skipif(not os.getenv("EKP_SOURCE_PATH"),reason="Set EKP_SOURCE_PATH to validate the user's original workbook")
def test_actual_classifier_all_cells_and_mapping():
    data=Path(os.environ["EKP_SOURCE_PATH"]).read_bytes()
    parsed=parse_workbook(data,"искл_пожар_задымление.xlsx")
    stats=parsed["metadata"]["statistics"]
    assert not parsed["validation_errors"]
    assert (stats["source_rows"],stats["incident_count"],stats["group_count"])==(1305,1281,24)
    assert stats["service_count"]==53 and stats["rule_count"]==20913
    assert stats["no_response_rule_count"]==741
    wb=load_workbook(io.BytesIO(data),data_only=True)
    for row in parsed["rows"]:
        assert len(row["raw"])==90
        for column,value in row["raw"].items():
            assert value==wb[row["source_sheet"]][f'{column}{row["source_row"]}'].value
    wb.close()
    incidents={r["code"]:r for r in parsed["rows"] if r["kind"]=="incident"}
    assert incidents["2010000"]["title"]=="ДТП без пострадавших"
    assert incidents["2010000"]["response_scenario_code"]=="2_4"
    assert incidents["24120100"]["response_scenario_code"] is None
    assert incidents["1010101"]["title"]=="пожар: мусор"
    assert any(rule["column"]=="AY" for row in incidents.values() for rule in row["rules"])
    assert any(rule["column"]=="CL" for row in incidents.values() for rule in row["rules"])
